from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import fitz  # pymupdf
from docx import Document as DocxDocument
from docx.table import Table

from .ocr import ocr_tesseract

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
# a page with less text than this is treated as a scan and sent to OCR
MIN_TEXT_CHARS = 25
OCR_DPI = 200  # lower gets shaky on small print, higher is slow with little gain


@dataclass
class Page:
    number: int
    text: str
    method: str  # "text", "ocr" or "docx"


@dataclass
class Document:
    path: Path
    pages: list[Page]

    @property
    def ocr_pages(self) -> int:
        return sum(p.method == "ocr" for p in self.pages)


def extract(path: Path, ocr: Callable[[bytes], str] = ocr_tesseract) -> Document:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return Document(path, _pdf_pages(path, ocr))
    if suffix == ".docx":
        # Word files have no fixed pages, so the whole file counts as page 1
        return Document(path, [Page(1, _docx_text(path), "docx")])
    if suffix in {".txt", ".md"}:
        return Document(path, [Page(1, path.read_text(encoding="utf-8", errors="replace"), "text")])
    raise ValueError(f"unsupported file type: {suffix or path.name}")


def _pdf_pages(path: Path, ocr: Callable[[bytes], str]) -> list[Page]:
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            text = page.get_text().strip()
            if len(text) >= MIN_TEXT_CHARS:
                pages.append(Page(i + 1, text, "text"))
                continue
            png = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
            pages.append(Page(i + 1, ocr(png).strip(), "ocr"))
    return pages


def _docx_text(path: Path) -> str:
    parts = []
    # iter_inner_content keeps paragraphs and tables in the order they appear in the file
    for item in DocxDocument(str(path)).iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat, so dedupe them
        else:
            parts.append(item.text)
    return "\n".join(parts)
