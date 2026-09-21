import docx
import pytest
from conftest import needs_tesseract, scanned_pdf, text_pdf

from docingest.extract import MIN_TEXT_CHARS, extract
from docingest.ocr import OCRError, ocr_tesseract, ocr_vlm


def never_ocr(png):
    raise AssertionError("OCR was called for a page that has text")


def test_pdf_with_a_text_layer_is_read_directly(tmp_path):
    text_pdf(tmp_path / "a.pdf", ["First page has plenty of words on it.", "Second page also has plenty of words."])
    doc = extract(tmp_path / "a.pdf", ocr=never_ocr)
    assert [(p.number, p.method) for p in doc.pages] == [(1, "text"), (2, "text")]
    assert "Second page" in doc.pages[1].text and doc.ocr_pages == 0


def test_a_page_with_almost_no_text_is_sent_to_ocr(tmp_path):
    scanned_pdf(tmp_path / "scan.pdf", "hello")
    seen = []

    def fake_ocr(png):
        seen.append(png[:8])
        return "  recognised words  "

    doc = extract(tmp_path / "scan.pdf", ocr=fake_ocr)
    assert seen == [b"\x89PNG\r\n\x1a\n"]  # it was handed a real PNG
    assert (doc.pages[0].method, doc.pages[0].text, doc.ocr_pages) == ("ocr", "recognised words", 1)


def test_mixed_pdf_only_ocrs_the_scanned_pages(tmp_path):
    import fitz

    scanned_pdf(tmp_path / "scan.pdf", "scanned")
    text_pdf(tmp_path / "text.pdf", ["A page that has a proper text layer with enough characters."])
    merged = fitz.open()
    for name in ("text.pdf", "scan.pdf"):
        with fitz.open(tmp_path / name) as part:
            merged.insert_pdf(part)
    merged.save(tmp_path / "mixed.pdf")
    doc = extract(tmp_path / "mixed.pdf", ocr=lambda png: "from ocr")
    assert [p.method for p in doc.pages] == ["text", "ocr"]


def test_threshold_is_about_characters_not_pages():
    assert MIN_TEXT_CHARS > 5  # a stray page number must not count as a page with text


def test_docx_keeps_paragraphs_and_tables_in_order(tmp_path):
    d = docx.Document()
    d.add_paragraph("Before the table")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([("Item", "Price"), ("Tea", "3")]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_paragraph("After the table")
    d.save(tmp_path / "f.docx")

    text = extract(tmp_path / "f.docx").pages[0].text
    assert text.splitlines() == ["Before the table", "Item | Price", "Tea | 3", "After the table"]


def test_plain_text_and_unsupported_types(tmp_path):
    (tmp_path / "n.txt").write_text("plain")
    assert extract(tmp_path / "n.txt").pages[0].text == "plain"
    (tmp_path / "x.xlsx").write_bytes(b"")
    with pytest.raises(ValueError, match="unsupported"):
        extract(tmp_path / "x.xlsx")


@needs_tesseract
def test_tesseract_actually_reads_a_scanned_page(tmp_path):
    scanned_pdf(tmp_path / "scan.pdf", "INVOICE TOTAL 450 DUE FRIDAY")
    doc = extract(tmp_path / "scan.pdf")
    assert doc.pages[0].method == "ocr"
    text = doc.pages[0].text.upper()
    assert "INVOICE" in text and "450" in text


def test_missing_tesseract_is_a_clear_error(monkeypatch):
    monkeypatch.setenv("PATH", "")
    with pytest.raises(OCRError, match="brew install tesseract"):
        ocr_tesseract(b"not an image")


class FakeLLM:
    def chat(self, messages, **kw):
        self.messages, self.kw = messages, kw
        return {"content": "transcribed"}


def test_vlm_ocr_sends_the_image_inline():
    llm = FakeLLM()
    assert ocr_vlm(b"\x89PNGdata", llm) == "transcribed"
    parts = llm.messages[0]["content"]
    assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
    assert "model" not in llm.kw  # left to the client's own model


def test_vlm_ocr_can_use_a_separate_vision_model():
    llm = FakeLLM()
    ocr_vlm(b"\x89PNGdata", llm, model="my-vision-model")
    assert llm.kw["model"] == "my-vision-model"
